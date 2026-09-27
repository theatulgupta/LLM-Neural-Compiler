"""Build the ten mentor progress reports as editable Word files and PDFs."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "progress"
FIG = OUT / "figures"
FONT = "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"
SANS = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
SANS_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

INK = (26, 29, 33)
MUTED = (90, 98, 108)
LINE = (197, 205, 214)
GREEN = (11, 107, 79)
GREEN_BG = (232, 243, 238)
RED = (155, 28, 28)
RED_BG = (250, 236, 236)
PAPER = (255, 255, 255)

AUTHOR = "Atul Kumar Gupta"
ROLL = "25MCSS06"
TITLE = "LLM-Guided Neural Network Compilation for Real-time UAV Edge Inference"


def _font(bold: bool, size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(SANS_BOLD if bold else SANS, size)


def _text_size(draw: ImageDraw.ImageDraw, text: str, font: ImageFont.FreeTypeFont) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def _center(draw: ImageDraw.ImageDraw, text: str, xy: tuple[int, int, int, int], font, fill) -> None:
    x0, y0, x1, y1 = xy
    tw, th = _text_size(draw, text, font)
    draw.text((x0 + (x1 - x0 - tw) / 2, y0 + (y1 - y0 - th) / 2), text, font=font, fill=fill)


def _box(draw, xy, text, fill=PAPER, outline=INK, size=22) -> None:
    draw.rounded_rectangle(xy, radius=12, fill=fill, outline=outline, width=3)
    _center(draw, text, xy, _font(True, size), INK)


def _arrow(draw, x0, y, x1) -> None:
    draw.line((x0, y, x1 - 14, y), fill=INK, width=3)
    draw.polygon([(x1 - 16, y - 7), (x1, y), (x1 - 16, y + 7)], fill=INK)


def _save(name: str, image: Image.Image) -> Path:
    FIG.mkdir(parents=True, exist_ok=True)
    path = FIG / name
    image.save(path, "PNG")
    return path


def fig_split() -> Path:
    image = Image.new("RGB", (1100, 1400), PAPER)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((80, 80, 1020, 560), radius=16, fill=GREEN_BG, outline=GREEN, width=5)
    _center(draw, "Flight controller", (80, 220, 1020, 320), _font(True, 40), INK)
    _center(draw, "PX4 keeps the aircraft flying", (80, 330, 1020, 410), _font(False, 28), MUTED)
    draw.line((550, 580, 550, 700), fill=INK, width=4)
    draw.polygon([(530, 690), (570, 690), (550, 720)], fill=INK)
    draw.rounded_rectangle((80, 760, 1020, 1240), radius=16, fill=GREEN_BG, outline=GREEN, width=5)
    _center(draw, "Companion computer", (80, 900, 1020, 1000), _font(True, 40), INK)
    _center(draw, "The camera network runs here", (80, 1020, 1020, 1100), _font(False, 28), MUTED)
    return _save("01_split.png", image)


def fig_rule() -> Path:
    image = Image.new("RGB", (1100, 1400), PAPER)
    draw = ImageDraw.Draw(image)
    labels = [
        ("Suggest", "name a plan"),
        ("Check", "accept or cut it"),
        ("Rewrite", "change the ONNX graph"),
        ("Measure", "time and numeric error"),
    ]
    y = 60
    for index, (title, subtitle) in enumerate(labels):
        draw.rounded_rectangle((120, y, 980, y + 240), radius=16, fill=GREEN_BG, outline=GREEN, width=5)
        _center(draw, title, (120, y + 40, 980, y + 130), _font(True, 40), INK)
        _center(draw, subtitle, (120, y + 130, 980, y + 200), _font(False, 26), MUTED)
        if index < 3:
            mid = y + 250
            draw.line((550, y + 246, 550, mid + 20), fill=INK, width=4)
            draw.polygon([(530, mid + 8), (570, mid + 8), (550, mid + 32)], fill=INK)
        y += 330
    return _save("02_rule.png", image)


def fig_fuse() -> Path:
    image = Image.new("RGB", (1100, 1300), PAPER)
    draw = ImageDraw.Draw(image)
    draw.text((80, 36), "Before", font=_font(True, 28), fill=MUTED)
    y = 90
    for label in ("Conv", "Relu", "Flatten 64", "Gemm K = 64"):
        draw.rounded_rectangle((80, y, 1020, y + 90), radius=12, fill=GREEN_BG, outline=GREEN, width=4)
        _center(draw, label, (80, y, 1020, y + 90), _font(True, 28), INK)
        y += 105
    draw.text((80, y + 8), "After graph_fuse", font=_font(True, 28), fill=MUTED)
    y += 60
    for label in ("FusedConv", "Flatten 64", "Gemm K = 64"):
        draw.rounded_rectangle((80, y, 1020, y + 90), radius=12, fill=GREEN_BG, outline=GREEN, width=4)
        _center(draw, label, (80, y, 1020, y + 90), _font(True, 28), INK)
        y += 105
    return _save("03_fuse.png", image)


def fig_zoo() -> Path:
    image = Image.new("RGB", (1100, 1400), PAPER)
    draw = ImageDraw.Draw(image)
    names = [
        "yolov8n    people and vehicles",
        "yolo11n    people and vehicles",
        "yolov8n-pose    keypoints",
        "yolov8n-seg    masks",
        "ssdlite    lighter detector",
        "mobilenet    classifier",
        "midas    depth",
    ]
    y = 40
    for name in names:
        draw.rounded_rectangle((70, y, 1030, y + 100), radius=12, fill=GREEN_BG, outline=GREEN, width=4)
        _center(draw, name, (70, y, 1030, y + 100), _font(True, 24), INK)
        y += 120
    draw.line((550, y - 10, 550, y + 30), fill=INK, width=4)
    draw.polygon([(530, y + 20), (570, y + 20), (550, y + 44)], fill=INK)
    draw.rounded_rectangle((180, y + 50, 920, y + 170), radius=14, fill=INK, outline=INK, width=3)
    _center(draw, "one compile loop", (180, y + 50, 920, y + 170), _font(True, 32), PAPER)
    return _save("04_zoo.png", image)


def fig_gate() -> Path:
    image = Image.new("RGB", (1100, 1200), PAPER)
    draw = ImageDraw.Draw(image)
    y = 40
    for index in range(3):
        draw.rounded_rectangle((160, y, 940, y + 160), radius=14, fill=GREEN_BG, outline=GREEN, width=4)
        _center(draw, f"Try {index + 1}", (160, y, 940, y + 160), _font(True, 36), INK)
        draw.line((550, y + 166, 550, y + 210), fill=INK, width=4)
        draw.polygon([(530, y + 200), (570, y + 200), (550, y + 224)], fill=INK)
        y += 250
    draw.rounded_rectangle((120, y, 980, y + 150), radius=14, fill=GREEN_BG, outline=GREEN, width=4)
    _center(draw, "Accepted plan", (120, y, 980, y + 150), _font(True, 32), INK)
    y += 200
    draw.rounded_rectangle((120, y, 980, y + 150), radius=14, fill=RED_BG, outline=RED, width=4)
    _center(draw, "Heuristic fallback", (120, y, 980, y + 150), _font(True, 32), INK)
    return _save("05_gate.png", image)


def fig_bars() -> Path:
    image = Image.new("RGB", (1100, 900), PAPER)
    draw = ImageDraw.Draw(image)
    rows = [
        (700, "preset", GREEN),
        (460, "chosen", GREEN),
        (300, "faster, failed the check", RED),
    ]
    y = 80
    for width, label, color in rows:
        draw.text((60, y), label, font=_font(True, 26), fill=INK)
        draw.rectangle((60, y + 50, 60 + width, y + 150), fill=color)
        y += 230
    draw.text((60, 800), "Shorter bar means lower p50. Red is not chosen.", font=_font(False, 22), fill=MUTED)
    return _save("06_bars.png", image)


def fig_planes() -> Path:
    image = Image.new("RGB", (1100, 1300), PAPER)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((80, 60, 1020, 360), radius=16, fill=GREEN_BG, outline=GREEN, width=5)
    _center(draw, "Compile", (80, 140, 1020, 230), _font(True, 40), INK)
    _center(draw, "plan, rewrite, measure", (80, 230, 1020, 310), _font(False, 26), MUTED)
    draw.line((550, 370, 550, 470), fill=INK, width=4)
    draw.polygon([(530, 458), (570, 458), (550, 486)], fill=INK)
    draw.rounded_rectangle((180, 500, 920, 700), radius=16, fill=INK, outline=INK, width=3)
    _center(draw, "ONNX file + sidecar", (180, 500, 920, 700), _font(True, 32), PAPER)
    draw.line((550, 710, 550, 810), fill=INK, width=4)
    draw.polygon([(530, 798), (570, 798), (550, 826)], fill=INK)
    draw.rounded_rectangle((80, 840, 1020, 1140), radius=16, fill=GREEN_BG, outline=GREEN, width=5)
    _center(draw, "Deploy", (80, 920, 1020, 1010), _font(True, 40), INK)
    _center(draw, "Gazebo camera, ROS node", (80, 1010, 1020, 1090), _font(False, 26), MUTED)
    return _save("07_planes.png", image)


def fig_hosts() -> Path:
    image = Image.new("RGB", (1100, 1300), PAPER)
    draw = ImageDraw.Draw(image)
    hosts = ["Groq", "OpenAI", "Gemini", "DeepSeek", "Anthropic", "A local model"]
    y = 40
    for name in hosts:
        draw.rounded_rectangle((140, y, 960, y + 90), radius=12, fill=GREEN_BG, outline=GREEN, width=4)
        _center(draw, name, (140, y, 960, y + 90), _font(True, 28), INK)
        y += 110
    draw.line((550, y - 10, 550, y + 20), fill=INK, width=4)
    draw.polygon([(530, y + 10), (570, y + 10), (550, y + 34)], fill=INK)
    draw.rounded_rectangle((120, y + 40, 980, y + 180), radius=14, fill=INK, outline=INK, width=3)
    _center(draw, "one HTTP client", (120, y + 40, 980, y + 180), _font(True, 32), PAPER)
    return _save("08_hosts.png", image)


def fig_page() -> Path:
    image = Image.new("RGB", (1100, 1300), PAPER)
    draw = ImageDraw.Draw(image)
    bands = [
        "This machine, and whether a GPU is here",
        "What the graph allows",
        "The plan, and each attempt",
        "Which measured plan was kept",
        "The seven models, and a fresh timing",
    ]
    y = 50
    for band in bands:
        draw.rounded_rectangle((70, y, 1030, y + 190), radius=16, fill=GREEN_BG, outline=GREEN, width=4)
        _center(draw, band, (90, y, 1010, y + 190), _font(True, 26), INK)
        y += 240
    return _save("09_page.png", image)


def fig_claim() -> Path:
    image = Image.new("RGB", (1100, 1200), PAPER)
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((60, 40, 1040, 760), radius=16, fill=GREEN_BG, outline=GREEN, width=5)
    _center(draw, "Can say now", (60, 70, 1040, 150), _font(True, 32), INK)
    left = [
        "An allowlisted plan",
        "A real ONNX rewrite",
        "A measured choice on this CPU",
        "One extra try after measurement",
        "ROS only loads the file",
    ]
    y = 190
    for line in left:
        draw.text((110, y), line, font=_font(False, 28), fill=INK)
        y += 100
    draw.rounded_rectangle((60, 820, 1040, 1120), radius=16, fill=RED_BG, outline=RED, width=5)
    _center(draw, "Waits for a board", (60, 860, 1040, 940), _font(True, 32), INK)
    _center(draw, "TensorRT, power, labelled accuracy", (60, 960, 1040, 1040), _font(False, 26), INK)
    return _save("10_claim.png", image)


Report = dict


def reports() -> list[Report]:
    return [
        {
            "num": 1,
            "slug": "companion",
            "heading": "Why the companion computer is a compile problem",
            "intro": [
                "A small drone splits the work in two. PX4, on the flight controller, keeps the aircraft in the air. A Linux computer beside it has to run the camera network, and it has to finish before the next frame is due.",
            ],
            "body": [
                "That second computer is the companion. It is where people, vehicles, a landing patch, or a depth cue would be recognised. The flight controller does not run those networks. If the companion is late, the aircraft is still flying, but it is flying without the picture.",
                "Turning a trained network into something the companion can run is a compile recipe. The recipe chooses which operations may be joined, whether the weights stay in 32-bit or drop to 8-bit, and how many CPU threads are used. In practice that recipe is still picked by habit. A choice that is fast on a desktop GPU is often slow on a four-core ARM CPU, and an 8-bit version can be both faster and numerically wrong.",
                "The lab tools around that computer had to be named before any compiler was written. PX4 is the flight stack, in the air and in simulation. Gazebo is the simulated world. ROS 2 is the message bus: programs publish and subscribe on topics, they do not call each other directly. ONNX is the file that holds a trained network as a graph of operations and tensors, which is the object a compiler can inspect.",
                "A camera that produces a frame about every 33 ms does not leave a large budget. If the network and the recipe around it take longer than that, frames are dropped. The recipe is therefore part of the research problem, not a setup step that can be left to taste.",
                "The first sketch treated the drone as one AI program. That was the wrong cut. PX4 already flies. Putting the camera network on the flight controller would mix a vision bug with a control bug, and a compile result would have nowhere honest to run. So the early work was the stack itself: a simulated aircraft, a camera, and a message bus, before any optimizer existed. The principle that survived is small. The science is the recipe on the companion. It is not a new autopilot.",
            ],
            "bullets": [],
            "caption": "Figure 1. PX4 flies. The companion computer is the only place the camera network runs.",
            "figure": fig_split,
            "outro": "Flight and vision are separate. What is still open is who is allowed to choose the compile recipe.",
        },
        {
            "num": 2,
            "slug": "rule",
            "heading": "The language model may only name a plan",
            "intro": [
                "The companion computer still needs a compile recipe. The question here is who may choose it.",
            ],
            "body": [
                "Asked to make a network fast, a language model will name passes that do not exist, and it will quote a speed it never measured. Either mistake is enough to make the result unusable in a thesis. A quoted frame rate that was not timed is not a result. A pass with no implementation is not a compiler.",
                "The model is therefore limited to one job. It may name a plan, and the plan must be taken from a fixed list. It may not edit the ONNX file. It may not fly the aircraft. Three other parts do the rest. A checker accepts the plan or cuts it. An engine rewrites the graph. A profiler measures time and compares the new outputs with the untouched network.",
                "The plan arrives as JSON and is checked before any rewrite. If the reply contains a made-up speed, that phrase is removed. Every stored run records that no frame rate was claimed. The separation is the research claim of this step: the model proposes, and only the measured run is evidence.",
                "The tempting design was to let the model rewrite ONNX, or to let it pick TensorRT settings in free text. A trial prompt did what free text always does: it named a SiLU fusion that this CPU cannot run, and it offered a frame rate. That single trial killed the free-text design. The fixed list is not a lack of ambition. It is the only way a later table can say what was tried.",
            ],
            "bullets": [],
            "caption": "Figure 1. Suggest, check, rewrite, then measure. The model stops at the first box.",
            "figure": fig_rule,
            "outro": "The model is the planner. Whether the rewrite actually changes a graph is a separate question.",
        },
        {
            "num": 3,
            "slug": "fixture",
            "heading": "A tiny network proves the rewrite",
            "intro": [
                "The planner is not allowed to edit the graph. This report checks that the engine, which is allowed to edit it, really does.",
            ],
            "body": [
                "The network used here is small enough to read by hand. Its input is one map of size 1×1×8×8. A convolution keeps a spatial map, a ReLU follows it, a pooling layer shrinks the map, Flatten turns the result into a vector of 64 values, and a Gemm produces the output. The Gemm width has to be 64, because that is the Flatten width. A width of 16 was an early mistake on this machine. It is kept as a test that must fail, so a broken contract cannot be reported as a successful compile.",
                "The pass used is graph_fuse. It joins a convolution with an immediately following ReLU. On ONNX Runtime for the CPU, the ReLU node leaves the graph and a FusedConv node takes its place. The Flatten width is unchanged. The proof is the operator list before and after the pass, not a log line that merely says the pass ran.",
                "Tests lock both facts. The broken width must not compile as a valid model. graph_fuse on the correct fixture must leave one FusedConv and no ReLU. That is the first evidence that the rewrite is a real graph change.",
                "The width of 16 came from a real mistake, not from a made-up test. The first reading was that Gemm K should match the channel count, or a 4 by 4 spatial guess. The Flatten after pooling is 4 by 4 by 4, so K is 64. The session compiled only after that was corrected. Keeping the broken file means the same wrong guess cannot pass quietly later. A log line that says fusion ran was also not enough. The operator list has to change, or the pass is only a name.",
            ],
            "bullets": [],
            "caption": "Figure 1. Conv followed by ReLU becomes one FusedConv. Flatten stays at 64 values.",
            "figure": fig_fuse,
            "outro": "The rewrite is visible on this small graph. The same commands still have to be shown on the networks a companion computer would actually run.",
        },
        {
            "num": 4,
            "slug": "zoo",
            "heading": "One loop, seven companion models",
            "intro": [
                "The rewrite was shown on a toy graph. The same compile commands are used here on the networks a companion computer would run.",
            ],
            "body": [
                "There is no second compiler for YOLO. Each network is one record: a task, an ONNX file, and an export script. The pipeline never branches on the model name.",
            ],
            "bullets": [
                "yolov8n and yolo11n detect people and vehicles.",
                "yolov8n-pose detects body keypoints.",
                "yolov8n-seg produces masks.",
                "ssdlite_mobilenetv3 is a lighter detector.",
                "mobilenetv3_small is a small classifier.",
                "midas_small estimates depth from one camera.",
            ],
            "after": [
                "Ultralytics YOLO graphs use SiLU, which is Conv, then Sigmoid, then Mul. The CPU fusion pass only joins Conv and Relu, so those YOLO graphs do not get smaller. MobileNet and MiDaS do get smaller where Conv is followed by batch-norm or Relu. That difference is a result.",
                "There is no fused CPU kernel for SiLU in this project, and none will be invented to make the node count fall. The list of models lives in experiments/zoo.yaml. Adding a model means a new record and an export script, not a new branch in the compiler.",
                "The first design was a YOLO compiler. That would have made every new camera task a new program. The node counts then looked like a failed fusion: several YOLO graphs stayed the same size. The pass was fine. Those graphs do not contain the Conv-ReLU pair it knows how to join. Writing a pretend SiLU kernel would have made the table look better and the system dishonest. MobileNet and MiDaS shrinking, on the same pass, is what showed the pass was alive.",
            ],
            "caption": "Figure 1. Seven models enter one compile loop.",
            "figure": fig_zoo,
            "outro": "One loop covers these seven models. A plan can still be refused before that loop rewrites anything.",
        },
        {
            "num": 5,
            "slug": "gate",
            "heading": "A plan can be refused",
            "intro": [
                "The seven models can be compiled. This report is only the gate that sits in front of that rewrite.",
            ],
            "body": [
                "A plan is JSON. It holds at most eight steps. Each step name must come from a closed list. The plan may also set how hard ONNX Runtime optimizes the graph, how many threads to use, and whether execution is sequential or parallel.",
                "An unknown name rejects the whole plan. The engine never sees it. A known name that does not match this graph, such as batch-norm fusion when the graph has no batch norm, is dropped. The rest of the plan continues.",
                "The checker allows three tries. If the reply is still illegal, an offline heuristic plan is used. That result is labelled as a fallback.",
                "On this CPU, a thread count above the number of cores is cut down to that number. Parallel execution is left as the plan asked for it. The checker does not rename the plan. It only removes or clamps what cannot run.",
                "Crashing on a useless step was the first checker. It made ordinary YOLO plans look broken, because a batch-norm fusion is legal in the list and absent in those graphs. A crash cannot be compared with another plan. A drop can. An unknown name is different: that is the model leaving the list, so the whole plan is refused. After three bad replies the offline plan is used, and it has to be labelled as a fallback. An earlier table called an 8-bit heuristic plan by a fusion name. The name lied, and the table could not be read. The name now has to match the steps.",
            ],
            "bullets": [],
            "caption": "Figure 1. Three tries. Then an accepted plan, or the heuristic fallback.",
            "figure": fig_gate,
            "outro": "An illegal plan is rejected or cut before the engine runs. Which of the legal plans is kept is decided only after they are measured.",
        },
        {
            "num": 6,
            "slug": "winner",
            "heading": "The winner is the fastest plan that passed",
            "intro": [
                "A plan can be legal and still be a bad compile. This report is about which legal plan is kept.",
            ],
            "body": [
                "Several legal plans are compiled on the same machine: fixed presets, the heuristic, and the model's plan. Each run records whether compile succeeded, whether the outputs still match the untouched graph, and the latency percentile p50.",
                "The chosen plan is the one with the lowest p50 among the plans that passed the numeric check. A plan that is faster but fails the check stays in the table. It is not chosen. If the model's plan has no rank, that is the reason, not a missing run.",
                "p50 is the median latency of repeated runs on this machine. It moves when the QEMU load moves, so two runs from different sessions are not one experiment. The numeric check compares the rewritten network with the original 32-bit graph.",
                "Each row stores where the plan came from, its name, the node count before and after the rewrite, whether the graph changed, and whether the numeric check passed. The model's row is never deleted when it matches a preset. It is marked as the same plan, so the record still shows what the model said.",
                "Sorting by latency alone was the obvious rule, and it is the wrong one. Eight-bit plans were often the fastest and they were also the ones that failed the comparison with the original 32-bit graph. A detector that is quicker and no longer agrees with itself is not a result. On one segmentation model the revised plan was faster and failed the check, so it was not chosen. On a detection model the model's second try got closer and still lost to a preset. Both outcomes stay in the record. Hiding them would turn the table into a advertisement for the model.",
            ],
            "bullets": [],
            "caption": "Figure 1. The short green bar can win. The red bar is faster and still rejected.",
            "figure": fig_bars,
            "outro": "A faster plan that fails the numeric check is not the result. The file that is kept still has to be loaded somewhere.",
        },
        {
            "num": 7,
            "slug": "deploy",
            "heading": "The aircraft only loads the file",
            "intro": [
                "Measurement ends with an ONNX file and a sidecar of session options. This report is the program that loads that file on the simulated aircraft.",
            ],
            "body": [
                "Gazebo publishes the camera on the ROS topic /nnc/camera/image_raw. The simulated airframe is nnc_x500_cam. PX4 SITL flies beside it. The ROS node inference_node loads the artifact and runs it on those frames.",
                "The node does not send flight setpoints. The compiler does not read the flight controller. The aircraft is a place to run the file.",
                "The sidecar next to the ONNX file stores the plan and the session options, so the node uses the same thread count that was measured. Camera frames can be written to disk with record_frames.py. That recording is a separate experiment from the compile table.",
                "The early write-up treated the flying stack as the thesis, so the compiler was going to talk to PX4. That would have made a bad setpoint look like a bad compile. The node was therefore forbidden from publishing setpoints, and the compiler was forbidden from reading them. Getting a picture out of Gazebo was its own failure: one render path crashed, and the camera topic was empty until the bridge and the airframe matched. Those are environment facts. They are not compile results, and they are not written down as if they were.",
            ],
            "bullets": [],
            "caption": "Figure 1. The compile side and the aircraft share only the artifact.",
            "figure": fig_planes,
            "outro": "The simulated aircraft only runs the file. It does not choose the recipe. A remote model can still be the thing that names the plan.",
        },
        {
            "num": 8,
            "slug": "hosted",
            "heading": "A hosted model, one revision",
            "intro": [
                "The plan already has to pass a gate, and only a measured plan is kept. The planner itself can be a model on a remote host.",
            ],
            "body": [
                "There is one HTTP client. Groq is one named host. OpenAI, Anthropic, Gemini, DeepSeek, and other hosts that speak the same chat API are further rows. A new host is a URL and a model id. The API key stays outside git.",
                "After a plan has been measured, there is at most one extra call. If compile or the numeric check failed, that call is named llm_revised. If the plan passed but was strictly slower than the best legal plan, the call is named llm_improved. A rate limit is stored as the text LLM HTTP 429. The extra plan is allowed to lose.",
                "A slow reasoning model needs a longer timeout and sometimes no temperature field. Those are settings in the environment, not a new program. The same three tries from the gate still apply on that extra call. Only one extra compile is run.",
                "Groq was hard-coded at first. A second host would have meant a second program, and a rate limit would have looked like a missing result. Pose hit HTTP 429 on the first call. The row is the heuristic, and the reason is stored. MobileNet hit 429 on the extra try. That compile was skipped instead of saving the heuristic under the name of a revision. The extra try itself used to die on one bad reply. It now gets the same three tries as the first plan, and it still may compile only once. The model is allowed to lose. That is the point of writing the rank down.",
            ],
            "bullets": [],
            "caption": "Figure 1. Many hosts, one client, then one extra measured try.",
            "figure": fig_hosts,
            "outro": "A remote model can name a plan, and the record can show that the plan lost. Running that loop by hand is still a long sequence of commands.",
        },
        {
            "num": 9,
            "slug": "page",
            "heading": "A page that runs the loop",
            "intro": [
                "Compile, check, and measure were separate commands. They are the same functions here, called from one page on this machine.",
            ],
            "body": [
                "The page shows the host, the graph notes, the plan and its attempts, the ranking, a table across the seven models, and a fresh timing of the chosen file. The operator picks a host and a model id. If the key is missing, the page says key not set. It does not ask for the key.",
                "The seven ONNX files are stored in experiments/models.",
                "The page listens only on this machine, at 127.0.0.1. One job runs at a time. If the console process is down, the buttons become usable again and the bar says the console is not reachable. Analyze, plan, optimize, export, and a remeasure of the chosen file all call the same Python functions as the command line.",
                "A markdown table was the first way to show the work. It hid the attempts. A refused step, a rate limit, and a faster plan that failed the check all lived in JSON that a reader would not open. The page had to show those, or the demonstration would slide back to the fastest number. An early version also locked every button after one failed request, which looked like a dead program. The lock now clears, and the bar says why.",
            ],
            "bullets": [],
            "caption": "Figure 1. The page is those five bands, in that order.",
            "figure": fig_page,
            "outro": "The loop can be run and shown without reading raw logs. What the thesis is allowed to claim is a separate statement.",
        },
        {
            "num": 10,
            "slug": "claim",
            "heading": "What can be claimed",
            "intro": [
                "The path runs from a companion computer that must see, to a measured file on a simulated camera. This report states the claim. It does not add another tool.",
            ],
            "body": [
                "The work can say that a language model names an allowlisted plan, that a checker can refuse it, that an engine rewrites ONNX, that a profiler measures, that the kept plan is the fastest one that passed, that one extra try is allowed after measurement, and that ROS only loads the file.",
                "It cannot yet say TensorRT speed, board power, or accuracy on labelled images. Those need a board with an NVIDIA GPU. This machine is QEMU on ARM and has none. Latency from this machine must not be compared with a Jetson.",
                "Every stored run sets fps_claimed to false. When a throughput number appears, it is 1000 divided by the mean latency in milliseconds. It is not a camera frame rate, and it is not a claim about a Jetson.",
                "TensorRT was the original target, because that is the usual edge compiler in the UAV papers. This machine is a QEMU ARM guest with a virtual GPU and no NVIDIA device. A TensorRT number here would have been invented. The same discipline applies to power and to accuracy on labelled frames. The claim stops where the evidence stops. The right column of the figure is not a failure of the method. It is the list of experiments that this host cannot run.",
            ],
            "bullets": [],
            "caption": "Figure 1. The left column is the thesis. The right column waits for a board.",
            "figure": fig_claim,
            "outro": "TensorRT speed and board power are the measurements that belong on a real board. They are not results of this machine.",
        },
    ]


def _set_run_font(run, name: str, size: int, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if color is not None:
        run.font.color.rgb = color


def write_docx(item: Report, figure: Path) -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Liberation Serif"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Liberation Serif")

    def add(text: str, *, size=11, bold=False, space_after=8, center=False) -> None:
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(space_after)
        para.paragraph_format.space_before = Pt(2)
        para.paragraph_format.line_spacing = 1.15
        if center:
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(text)
        _set_run_font(run, "Liberation Serif", 12 if size == 11 else size, bold)

    add("Thesis progress report", size=11, bold=True, space_after=2)
    add(TITLE, size=12, space_after=2)
    add(f"{AUTHOR}  ·  {ROLL}", size=11, space_after=10)
    add(item["heading"], size=16, bold=True, space_after=10)
    for paragraph in item["intro"]:
        add(paragraph)
    for paragraph in item["body"]:
        add(paragraph)
    for bullet in item["bullets"]:
        para = doc.add_paragraph(style="List Bullet")
        para.paragraph_format.space_after = Pt(2)
        run = para.add_run(bullet)
        _set_run_font(run, "Liberation Serif", 11)
    for paragraph in item.get("after", []):
        add(paragraph)
    doc.add_picture(str(figure), width=Inches(5.8))
    add(item["caption"], size=10, space_after=10)
    add(item["outro"])
    path = OUT / f"{item['num']:02d}_{item['slug']}.docx"
    doc.save(path)
    return path


class ReportPDF(FPDF):
    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("Serif", size=9)
        self.set_text_color(90, 98, 108)
        self.cell(0, 8, f"{AUTHOR}  ·  {ROLL}  ·  {self.page_no()}", align="C")


def write_pdf(item: Report, figure: Path) -> Path:
    pdf = ReportPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=16)
    pdf.add_font("Serif", "", FONT)
    pdf.add_font("Serif", "B", FONT_BOLD)
    pdf.add_page()
    pdf.set_text_color(26, 29, 33)

    def para(text: str, *, size: int = 12, bold: bool = False, height: float = 6.3) -> None:
        pdf.set_x(pdf.l_margin)
        pdf.set_font("Serif", "B" if bold else "", size)
        pdf.multi_cell(pdf.epw, height, text)
        pdf.ln(2.2)

    para("Thesis progress report", bold=True, height=6)
    para(TITLE, height=6)
    para(f"{AUTHOR}  ·  {ROLL}", height=6)
    pdf.ln(1)
    para(item["heading"], size=16, bold=True, height=8)
    pdf.ln(1)

    for paragraph in item["intro"] + item["body"]:
        para(paragraph)
    for bullet in item["bullets"]:
        para(f"- {bullet}")
    for paragraph in item.get("after", []):
        para(paragraph)
    pdf.image(str(figure), w=150)
    pdf.ln(3)
    para(item["caption"], size=10)
    para(item["outro"])
    path = OUT / f"{item['num']:02d}_{item['slug']}.pdf"
    pdf.output(str(path))
    return path


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for item in reports():
        figure = item["figure"]()
        docx = write_docx(item, figure)
        pdf = write_pdf(item, figure)
        print(docx.name, pdf.name)


if __name__ == "__main__":
    main()
